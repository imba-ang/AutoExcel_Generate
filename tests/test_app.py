from __future__ import annotations

import io
import json
import sqlite3
import sys
import types
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app.config import BASE_DIR, Settings
from app.db import init_db, list_students, load_answers, save_submission
from app.main import create_app


def make_client(tmp_path: Path) -> TestClient:
    settings = Settings(
        data_dir=tmp_path,
        template_path=BASE_DIR / "app" / "assets" / "template.xlsx",
        teacher_password="teacher-test-password",
        session_secret="test-secret-that-is-longer-than-32-characters",
        public_base_url="https://forms.example.edu",
        cookie_secure=False,
        qr_base_url="https://forms.example.cn",
    )
    return TestClient(create_app(settings))


def login(client: TestClient) -> None:
    response = client.post("/teacher/login", data={"password": "teacher-test-password"})
    assert response.status_code == 200


def test_student_can_submit_load_and_overwrite(tmp_path: Path):
    with make_client(tmp_path) as client:
        home = client.get("/")
        assert home.status_code == 200
        for path in ("qidian", "po", "kuo", "shai"):
            assert f'/student/{path}' in home.text

        page = client.get("/student/po")
        assert page.status_code == 200
        assert "四问破题" in page.text
        assert 'data-cell="E10"' in page.text
        assert 'data-cell="C6"' not in page.text
        assert 'class="sheet-column">A</th>' in page.text
        assert 'class="sheet-row">1</th>' in page.text
        assert '>提交</button>' in page.text
        assert "提交破表" not in page.text
        assert 'class="fixed-cell"' in page.text
        assert 'class="cell-value fixed-text"' in page.text
        for section in ("qidian", "po", "kuo", "shai"):
            fixed_cells = [
                cell
                for row in client.app.state.excel_template.render_rows(section)
                for cell in row["cells"]
                if not cell["editable"] and cell["value"]
            ]
            assert fixed_cells
            assert all("text-align:center" in cell["style"] for cell in fixed_cells)
            assert all("vertical-align:middle" in cell["style"] for cell in fixed_cells)

        payload = {
            "student_name": "张三",
            "student_id": "20260001",
            "answers": {"E10": "第一次答案", "NOT_ALLOWED": "忽略"},
        }
        assert client.post("/api/submissions/po", json=payload).status_code == 200

        payload["answers"] = {"E10": "修改后的答案"}
        assert client.post("/api/submissions/po", json=payload).status_code == 200
        loaded = client.post("/api/submissions/po/load", json={**payload, "answers": {}})
        assert loaded.status_code == 200
        assert loaded.json()["answers"] == {"E10": "修改后的答案"}

        starting_page = client.get("/student/qidian")
        assert starting_page.status_code == 200
        assert "先留下第一反应" in starting_page.text
        assert 'data-cell="C6"' in starting_page.text
        assert 'data-cell="E10"' not in starting_page.text


def test_student_name_must_match_existing_id(tmp_path: Path):
    with make_client(tmp_path) as client:
        payload = {
            "student_name": "张三",
            "student_id": "20260001",
            "answers": {"E17": "答案"},
        }
        assert client.post("/api/submissions/kuo", json=payload).status_code == 200
        payload["student_name"] = "李四"
        assert client.post("/api/submissions/kuo", json=payload).status_code == 409


def test_shai_loads_candidate_plans_from_kuo(tmp_path: Path):
    with make_client(tmp_path) as client:
        identity = {"student_name": "王五", "student_id": "20260002"}
        client.post(
            "/api/submissions/kuo",
            json={**identity, "answers": {"C23": "校园闲置物品交换平台"}},
        )
        response = client.post(
            "/api/submissions/shai/load", json={**identity, "answers": {}}
        )
        assert response.status_code == 200
        assert response.json()["context_answers"]["C23"] == "校园闲置物品交换平台"


def test_teacher_views_and_downloads_original_template_shape(tmp_path: Path):
    with make_client(tmp_path) as client:
        client.post(
            "/api/submissions/po",
            json={
                "student_name": "赵六",
                "student_id": "20260003",
                "answers": {"C14": "如何减少食堂排队时间？"},
            },
        )
        assert client.get("/teacher", follow_redirects=False).status_code == 303
        login(client)
        dashboard = client.get("/teacher")
        assert "赵六" in dashboard.text
        assert "20260003" in dashboard.text

        download = client.get("/teacher/download/20260003")
        assert download.status_code == 200
        workbook = load_workbook(io.BytesIO(download.content), data_only=False)
        sheet = workbook["4·4·4创新方案生成表"]
        assert sheet.max_row == 42
        assert sheet.max_column == 7
        assert "A1:G2" in {str(item) for item in sheet.merged_cells.ranges}
        assert sheet["A9"].alignment.horizontal == "center"
        assert sheet["A9"].alignment.vertical == "center"
        assert sheet["A9"].alignment.indent == 0
        assert sheet["A9"].alignment.relativeIndent == 0
        assert sheet["C14"].value == "我真正值得解决的问题是：如何减少食堂排队时间？"
        assert str(sheet["A31"].value).startswith("=IF(")

        batch = client.get("/teacher/download-all")
        assert batch.status_code == 200
        with zipfile.ZipFile(io.BytesIO(batch.content)) as archive:
            assert any(name.endswith("_创新方案.xlsx") for name in archive.namelist())


def test_qrcode_uses_configured_domestic_mirror_url(tmp_path: Path):
    with make_client(tmp_path) as client:
        login(client)
        page = client.get("/teacher/qrcodes")
        assert page.text.count("下载二维码") == 4
        assert "https://forms.example.cn/student/qidian" in page.text
        assert "https://forms.example.cn/student/po" in page.text
        targets: list[str] = []

        def make_qr(target: str):
            targets.append(target)
            return Image.new("1", (128, 128), color=1)

        sys.modules["qrcode"] = types.SimpleNamespace(
            make=make_qr
        )
        image = client.get("/teacher/qrcode/po.png")
        assert image.status_code == 200
        assert image.headers["content-type"] == "image/png"
        assert targets == ["https://forms.example.cn/student/po"]


def test_twenty_students_can_save_concurrently(tmp_path: Path):
    db_path = tmp_path / "concurrent.db"
    init_db(db_path)

    def save(index: int) -> None:
        save_submission(
            db_path,
            student_id=f"2026{index:04d}",
            student_name=f"学生{index}",
            section="po",
            answers={"E10": f"答案{index}"},
        )

    with ThreadPoolExecutor(max_workers=20) as executor:
        list(executor.map(save, range(20)))

    assert len(list_students(db_path)) == 20


def test_legacy_database_is_upgraded_and_starting_answers_are_preserved(tmp_path: Path):
    db_path = tmp_path / "legacy.db"
    timestamp = "2026-09-20T12:00:00+08:00"
    with sqlite3.connect(db_path) as connection:
        connection.executescript(
            """
            CREATE TABLE students (
                student_id TEXT PRIMARY KEY,
                student_name TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE submissions (
                student_id TEXT NOT NULL,
                section TEXT NOT NULL CHECK(section IN ('po', 'kuo', 'shai')),
                answers_json TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (student_id, section),
                FOREIGN KEY (student_id) REFERENCES students(student_id) ON DELETE CASCADE
            );
            """
        )
        connection.execute(
            "INSERT INTO students VALUES (?, ?, ?, ?)",
            ("20260088", "旧记录", timestamp, timestamp),
        )
        connection.execute(
            "INSERT INTO submissions VALUES (?, 'po', ?, ?, ?)",
            (
                "20260088",
                json.dumps({"C6": "旧问题", "C7": "旧答案", "E10": "事实"}, ensure_ascii=False),
                timestamp,
                timestamp,
            ),
        )

    init_db(db_path)

    assert load_answers(db_path, "20260088", "qidian") == {
        "C6": "旧问题",
        "C7": "旧答案",
    }
    save_submission(db_path, "20260088", "旧记录", "qidian", {"C6": "新问题"})
    assert load_answers(db_path, "20260088", "qidian") == {"C6": "新问题"}

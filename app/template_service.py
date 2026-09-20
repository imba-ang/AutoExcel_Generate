from __future__ import annotations

import io
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Color
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class SectionSpec:
    slug: str
    label: str
    subtitle: str
    start_row: int
    end_row: int
    editable_cells: tuple[str, ...]


SECTIONS: dict[str, SectionSpec] = {
    "qidian": SectionSpec(
        slug="qidian",
        label="起点",
        subtitle="先留下第一反应",
        start_row=5,
        end_row=7,
        editable_cells=("C6", "C7"),
    ),
    "po": SectionSpec(
        slug="po",
        label="破",
        subtitle="四问破题",
        start_row=9,
        end_row=14,
        editable_cells=("E10", "E11", "E12", "E13", "C14"),
    ),
    "kuo": SectionSpec(
        slug="kuo",
        label="扩",
        subtitle="四维发散",
        start_row=16,
        end_row=26,
        editable_cells=("E17", "E18", "E19", "E20", "C23", "C24", "C25", "C26"),
    ),
    "shai": SectionSpec(
        slug="shai",
        label="筛",
        subtitle="四筛收敛",
        start_row=28,
        end_row=40,
        editable_cells=(
            "C31", "D31", "E31", "F31", "G31",
            "C32", "D32", "E32", "F32", "G32",
            "C33", "D33", "E33", "F33", "G33",
            "C34", "D34", "E34", "F34", "G34",
            "C36", "A39",
        ),
    ),
}

FORMULA_SOURCES = {
    "A31": ("C23", "方案A"),
    "A32": ("C24", "方案B"),
    "A33": ("C25", "方案C"),
    "A34": ("C26", "方案D"),
}

PROMPT_PREFIX_CELLS = {"C14", "C36"}


def normalize_color(color: Color | None, fallback: str | None = None) -> str | None:
    if not color:
        return fallback
    if color.type == "rgb" and color.rgb:
        rgb = str(color.rgb)
        return f"#{rgb[-6:]}"
    return fallback


def side_css(name: str, side: Any) -> str | None:
    if not side or not side.style:
        return None
    widths = {"hair": 1, "thin": 1, "medium": 2, "thick": 3}
    styles = {"dashed": "dashed", "dotted": "dotted", "dashDot": "dashed"}
    width = widths.get(side.style, 1)
    line_style = styles.get(side.style, "solid")
    color = normalize_color(side.color, "#C7CBD1")
    return f"border-{name}:{width}px {line_style} {color}"


def cell_css(cell: Cell) -> str:
    parts: list[str] = []
    fill = normalize_color(cell.fill.fgColor)
    if fill and cell.fill.fill_type:
        parts.append(f"background-color:{fill}")

    font = cell.font
    if font.name:
        parts.append(f"font-family:{font.name},'Microsoft YaHei',sans-serif")
    if font.sz:
        parts.append(f"font-size:{font.sz}pt")
    if font.bold:
        parts.append("font-weight:700")
    if font.italic:
        parts.append("font-style:italic")
    font_color = normalize_color(font.color)
    if font_color:
        parts.append(f"color:{font_color}")

    alignment = cell.alignment
    if alignment.horizontal:
        horizontal = {"general": "left", "centerContinuous": "center"}.get(
            alignment.horizontal, alignment.horizontal
        )
        parts.append(f"text-align:{horizontal}")
    if alignment.vertical:
        vertical = {"center": "middle"}.get(alignment.vertical, alignment.vertical)
        parts.append(f"vertical-align:{vertical}")
    if alignment.wrap_text:
        parts.append("white-space:pre-wrap")

    for name in ("top", "right", "bottom", "left"):
        value = side_css(name, getattr(cell.border, name))
        if value:
            parts.append(value)
    return ";".join(parts)


class ExcelTemplate:
    def __init__(self, path: Path):
        self.path = path
        workbook = load_workbook(path, data_only=False)
        self.sheet_name = workbook.sheetnames[0]
        self.worksheet = workbook[self.sheet_name]

    @property
    def sections(self) -> list[SectionSpec]:
        return list(SECTIONS.values())

    def spec(self, slug: str) -> SectionSpec:
        if slug not in SECTIONS:
            raise KeyError(slug)
        return SECTIONS[slug]

    def _merge_for(self, row: int, column: int) -> tuple[Any | None, bool]:
        for merged in self.worksheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                is_anchor = row == merged.min_row and column == merged.min_col
                return merged, is_anchor
        return None, True

    def render_rows(
        self,
        slug: str,
        answers: dict[str, str] | None = None,
        context_answers: dict[str, str] | None = None,
        readonly: bool = False,
    ) -> list[dict[str, Any]]:
        spec = self.spec(slug)
        answers = answers or {}
        context_answers = context_answers or {}
        rows: list[dict[str, Any]] = []
        for row_number in range(spec.start_row, spec.end_row + 1):
            rendered_cells: list[dict[str, Any]] = []
            for column in range(1, self.worksheet.max_column + 1):
                merged, is_anchor = self._merge_for(row_number, column)
                if merged and not is_anchor:
                    continue
                cell = self.worksheet.cell(row_number, column)
                coordinate = cell.coordinate
                value = "" if cell.value is None else str(cell.value)
                formula_source = None
                if coordinate in FORMULA_SOURCES:
                    source_coordinate, fallback = FORMULA_SOURCES[coordinate]
                    formula_source = source_coordinate
                    value = context_answers.get(source_coordinate) or fallback
                editable = coordinate in spec.editable_cells
                prompt = value if editable and value else ""
                answer = answers.get(coordinate, "")
                display_value = answer if editable and answer else value
                if editable and coordinate in PROMPT_PREFIX_CELLS and readonly and answer:
                    display_value = f"{value}{answer}"
                style = cell_css(cell)
                if not editable and cell.value is not None:
                    style = (
                        f"{style};padding-left:0;padding-right:0;text-indent:0;"
                        "text-align:center;vertical-align:middle"
                    )
                rendered_cells.append(
                    {
                        "coordinate": coordinate,
                        "value": display_value,
                        "prompt": prompt,
                        "answer": answer,
                        "editable": editable and not readonly,
                        "input_cell": editable,
                        "formula_source": formula_source,
                        "colspan": (merged.max_col - merged.min_col + 1) if merged else 1,
                        "rowspan": (merged.max_row - merged.min_row + 1) if merged else 1,
                        "style": style,
                    }
                )
            height = self.worksheet.row_dimensions[row_number].height or 22
            rows.append(
                {
                    "number": row_number,
                    "display_number": row_number - spec.start_row + 1,
                    "height_px": round(height * 4 / 3),
                    "cells": rendered_cells,
                }
            )
        return rows

    def column_widths(self) -> list[int]:
        widths = []
        for column in range(1, self.worksheet.max_column + 1):
            letter = get_column_letter(column)
            width = self.worksheet.column_dimensions[letter].width or 13
            widths.append(round(width * 7 + 5))
        return widths

    def build_workbook(self, all_answers: dict[str, dict[str, str]]) -> io.BytesIO:
        workbook = load_workbook(self.path, data_only=False)
        worksheet = workbook[self.sheet_name]
        for slug, answers in all_answers.items():
            if slug not in SECTIONS:
                continue
            for coordinate, answer in answers.items():
                if coordinate not in SECTIONS[slug].editable_cells:
                    continue
                original = self.worksheet[coordinate].value or ""
                if coordinate in PROMPT_PREFIX_CELLS and answer:
                    worksheet[coordinate] = f"{original}{answer}"
                else:
                    worksheet[coordinate] = answer
        editable_coordinates = {
            coordinate
            for spec in SECTIONS.values()
            for coordinate in spec.editable_cells
        }
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.coordinate in editable_coordinates or cell.value is None:
                    continue
                alignment = copy(cell.alignment)
                alignment.horizontal = "center"
                alignment.vertical = "center"
                alignment.indent = 0
                alignment.relativeIndent = 0
                cell.alignment = alignment
        try:
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        except AttributeError:
            pass
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output
